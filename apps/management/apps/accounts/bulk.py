"""Bounded, previewed account batches. Never accept or generate password columns."""

import csv
import hashlib
import io
import json
import zipfile
from collections import Counter
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import serializers
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.audit.services import record_audit
from apps.tasks.services import enqueue
from common.errors import AppError, ErrorCode
from common.permissions import RequiresPasswordChange

from .administration import require_admin, serialize_admin_changes, snapshot
from .models import AccountBatch
from .serializers import UserCreateSerializer, UserSerializer, UserUpdateSerializer

IMPORT_FIELDS = [
    'username',
    'display_name',
    'email',
    'role_codes',
    'status',
    'node_ids',
    'quota_instances',
    'quota_vcpu',
    'quota_memory_mib',
    'quota_disk_gib',
    'remark',
]
EXPORT_FIELDS = IMPORT_FIELDS + ['created_at', 'updated_at', 'last_login']
SENSITIVE = {'email', 'remark'}
MAX_ROWS = 1000
MAX_BYTES = 5 * 1024 * 1024


def safe_cell(value):
    text = str(value if value is not None else '')
    return (
        "'" + text
        if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n'))
        else text
    )


def file_bytes(headers, rows, format):
    if format == 'xlsx':
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet('users')
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_cell(row.get(field, '')) for field in headers])
        target = io.BytesIO()
        workbook.save(target)
        return target.getvalue()
    output = io.StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows([safe_cell(row.get(field, '')) for field in headers] for row in rows)
    return ('\ufeff' + output.getvalue()).encode('utf-8')


def download(data, format, name):
    response = HttpResponse(
        data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        if format == 'xlsx'
        else 'text/csv; charset=utf-8',
    )
    response['Content-Disposition'] = f'attachment; filename="{name}.{format}"'
    response['Cache-Control'] = 'no-store'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


def parse_file(upload):
    if not upload or upload.size > MAX_BYTES:
        raise AppError(ErrorCode.INVALID_ARGUMENT, '请选择不超过 5 MB 的 CSV 或 XLSX 文件')
    data = upload.read(MAX_BYTES + 1)
    name = upload.name.lower()
    try:
        if name.endswith('.csv'):
            reader = csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''))
            values = []
            for row in reader:
                if len(values) > MAX_ROWS:
                    raise ValueError('每批最多 1000 行')
                values.append(row)
        elif name.endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024
                ):
                    raise ValueError('XLSX 解压内容超过限制')
                if any(x.flag_bits & 1 for x in archive.infolist()):
                    raise ValueError('不支持加密文件')
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise ValueError('仅支持单工作表文件')
                sheet = workbook.worksheets[0]
                if (sheet.max_column and sheet.max_column > len(IMPORT_FIELDS)) or (
                    sheet.max_row and sheet.max_row > MAX_ROWS + 1
                ):
                    raise ValueError('XLSX 行列数超限')
                values = []
                for row in sheet.iter_rows():
                    if len(values) > MAX_ROWS or len(row) > len(IMPORT_FIELDS):
                        raise ValueError('XLSX 行列数超限')
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('导入文件不允许公式')
                    values.append(
                        [str(cell.value) if cell.value is not None else '' for cell in row]
                    )
            finally:
                workbook.close()
        else:
            raise ValueError('仅支持 CSV 和 XLSX')
        if not values:
            raise ValueError('文件为空')
        headers = [str(value).strip() for value in values.pop(0)]
        if (
            not headers
            or len(headers) != len(set(headers))
            or 'username' not in headers
            or set(headers) - set(IMPORT_FIELDS)
        ):
            raise ValueError('表头缺失、重复或包含不允许的字段；请下载模板')
        rows = []
        for index, row in enumerate(values, 2):
            if name.endswith('.xlsx') and len(row) < len(headers):
                row += [''] * (len(headers) - len(row))
            if not any(str(v).strip() for v in row):
                continue
            if len(row) != len(headers):
                raise ValueError(f'第 {index} 行列数与表头不一致')
            if any(len(str(value)) > 2000 for value in row):
                raise ValueError(f'第 {index} 行字段过长')
            rows.append(
                {
                    'row': index,
                    'input': dict(zip(headers, [str(v).strip() for v in row], strict=True)),
                }
            )
        if not rows:
            raise ValueError('没有有效数据行')
        return rows
    except (ValueError, UnicodeError, zipfile.BadZipFile, csv.Error, KeyError) as error:
        raise AppError(
            ErrorCode.INVALID_ARGUMENT,
            str(error)
            if isinstance(error, ValueError) and not isinstance(error, UnicodeError)
            else '文件格式无效，CSV 请使用 UTF-8 编码',
        ) from error
    except Exception as error:  # noqa: BLE001 - library parser details must not expose uploaded content.
        raise AppError(
            ErrorCode.INVALID_ARGUMENT, '无法解析文件，请使用未加密的标准模板'
        ) from error


def prepare(row, mode):
    values = row['input']
    username = values.get('username', '').strip()
    matches = list(get_user_model().objects.filter(username__iexact=username)[:2])
    if len(matches) > 1:
        raise serializers.ValidationError({'username': '存在大小写冲突账号，请先由管理员处理'})
    user = matches[0] if matches else None
    if mode == 'create' and user:
        raise serializers.ValidationError({'username': '用户名已存在'})
    if mode == 'update' and not user:
        raise serializers.ValidationError({'username': '更新目标不存在'})
    data = {key: values[key] for key in ['email', 'remark', 'display_name'] if values.get(key)}
    for key in ['email', 'remark']:
        if data.get(key) == '__CLEAR__':
            data[key] = ''
    if values.get('role_codes'):
        data['roles'] = [v.strip() for v in values['role_codes'].split(';') if v.strip()]
    if values.get('node_ids'):
        raw = values['node_ids']
        if raw == '__CLEAR__':
            data['node_scope'] = None
        elif raw == '[]':
            data['node_scope'] = []
        else:
            try:
                data['node_scope'] = [int(v.strip()) for v in raw.split(';')]
            except ValueError as error:
                raise serializers.ValidationError({'node_ids': '节点 ID 使用分号分隔'}) from error
    quota = {}
    for field in ['instances', 'vcpu', 'memory_mib', 'disk_gib']:
        raw = values.get('quota_' + field, '')
        if raw:
            try:
                quota[field] = int(raw)
            except ValueError as error:
                raise serializers.ValidationError({'quota_' + field: '必须填写整数'}) from error
    if quota:
        data['quota_overrides'] = {
            **getattr(getattr(user, 'profile', None), 'quota_overrides', {}),
            **quota,
        }
    status = values.get('status', '')
    if status and status not in {'pending', 'active', 'disabled'}:
        raise serializers.ValidationError({'status': '仅支持 pending、active、disabled'})
    if user:
        if status == 'pending' and user.has_usable_password():
            raise serializers.ValidationError({'status': '不能将已初始化账号退回待激活'})
        if status:
            data['is_active'] = status == 'active'
        serializer = UserUpdateSerializer(user, data=data, partial=True)
    else:
        if not data.get('display_name') or not data.get('roles'):
            raise serializers.ValidationError({'detail': '新账号必须填写显示名和角色编码'})
        if status not in {'', 'pending'}:
            raise serializers.ValidationError(
                {'status': '新账号仅支持 pending；请重置密码后手动启用'}
            )
        data['username'] = username
        data['is_active'] = False
        serializer = UserCreateSerializer(data=data)
    serializer.is_valid(raise_exception=True)
    return user, serializer


def error_result(row, error):
    fields = (
        error.detail
        if isinstance(error, serializers.ValidationError)
        else error.details or {'detail': error.message}
    )
    return {
        'row': row['row'],
        'username': row['input'].get('username', ''),
        'status': 'ERROR',
        'code': getattr(error, 'code', 'VALIDATION_ERROR'),
        'errors': json.loads(json.dumps(fields, default=str)),
    }


def role_snapshot(serializer, user):
    from django.contrib.auth.models import Group

    from .role_views import describe_role

    codes = serializer.validated_data.get(
        'roles', list(user.groups.values_list('name', flat=True)) if user else ['Viewer']
    )
    roles = [
        {key: value for key, value in describe_role(group).items() if key != 'user_count'}
        for group in Group.objects.filter(name__in=codes).order_by('name')
    ]
    return hashlib.sha256(json.dumps(roles, sort_keys=True).encode()).hexdigest()


def preview(rows, mode):
    duplicates = Counter(row['input'].get('username', '').casefold() for row in rows)
    plan = []
    for row in rows:
        try:
            if duplicates[row['input'].get('username', '').casefold()] > 1:
                raise serializers.ValidationError(
                    {'username': '文件内用户名重复，所有重复行均不处理'}
                )
            user, serializer = prepare(row, mode)
            plan.append(
                {
                    **row,
                    'snapshot': snapshot(user),
                    'roles_snapshot': role_snapshot(serializer, user),
                    'status': 'UPDATE' if user else 'CREATE',
                    'target': user.pk if user else None,
                }
            )
        except (AppError, serializers.ValidationError) as error:
            plan.append({**row, **error_result(row, error)})
    return plan


def describe(batch):
    state = batch.state
    if batch.task_id and batch.task.status in {'CANCELLED', 'FAILED'}:
        state = batch.task.status
    response = {
        'id': str(batch.pk),
        'kind': batch.kind,
        'state': state,
        'mode': batch.mode,
        'format': batch.format,
        'strict': batch.strict,
        'task_id': str(batch.task_id) if batch.task_id else None,
        'expires_at': batch.expires_at,
        'total': len(batch.rows) if batch.kind == 'import' else len(batch.options.get('ids', [])),
        'results': batch.results,
        'download_ready': bool(batch.file_data) and batch.expires_at > timezone.now(),
    }
    if state == 'DRAFT':
        response['preview'] = [
            {
                key: value
                for key, value in row.items()
                if key not in {'input', 'snapshot', 'roles_snapshot'}
            }
            | {'values': row['input']}
            for row in batch.rows
        ]
    return response


def export_rows(users, fields):
    for user in users:
        profile = getattr(user, 'profile', None)
        quota = getattr(profile, 'quota_overrides', {})
        values = {
            'username': user.username,
            'display_name': getattr(profile, 'display_name', user.username),
            'email': user.email,
            'role_codes': ';'.join(sorted(user.groups.values_list('name', flat=True))),
            'status': UserSerializer(user).get_status(user),
            'node_ids': ';'.join(map(str, profile.node_scope))
            if profile and profile.node_scope
            else '[]'
            if profile and profile.node_scope == []
            else '',
            'remark': getattr(profile, 'remark', ''),
            'created_at': user.date_joined.isoformat(),
            'updated_at': profile.updated_at.isoformat() if profile else '',
            'last_login': user.last_login.isoformat() if user.last_login else '',
        }
        values.update({'quota_' + key: value for key, value in quota.items()})
        yield {field: values.get(field, '') for field in fields}


def execute_account_batch(task):
    batch = AccountBatch.objects.select_related('owner').get(pk=task.resource_id)
    required = 'accounts.import_user' if batch.kind == 'import' else 'accounts.export_user'
    require_admin(batch.owner, required)
    if batch.expires_at <= timezone.now():
        raise AppError(ErrorCode.CONFLICT, '批次已过期')
    if batch.state in {'SUCCEEDED', 'PARTIAL', 'FAILED'}:
        return {'batch_id': str(batch.pk), 'state': batch.state}
    batch.state = 'RUNNING'
    batch.save(update_fields=['state'])
    try:
        if batch.kind == 'export':
            fields = batch.options['fields']
            if SENSITIVE & set(fields):
                require_admin(batch.owner, 'accounts.export_sensitive_user')
            users = (
                get_user_model()
                .objects.filter(pk__in=batch.options['ids'])
                .select_related('profile')
                .prefetch_related('groups')
                .order_by('id')
            )
            batch.file_data = file_bytes(fields, export_rows(users, fields), batch.format)
            batch.state = 'SUCCEEDED'
            batch.results = [{'status': 'SUCCESS', 'count': users.count(), 'fields': fields}]
            batch.save()
        else:
            completed = {entry['row'] for entry in batch.results}
            if batch.strict:
                with transaction.atomic():
                    serialize_admin_changes()
                    current = preview(batch.rows, batch.mode)
                    if any(
                        row['status'] == 'ERROR'
                        or row.get('snapshot') != old.get('snapshot')
                        or row.get('roles_snapshot') != old.get('roles_snapshot')
                        for row, old in zip(current, batch.rows, strict=True)
                    ):
                        raise AppError(ErrorCode.CONFLICT, '严格模式重新校验失败，整批未执行')
                    process_rows(batch, completed, required)
            else:
                process_rows(batch, completed, required)
            batch.state = (
                'PARTIAL' if any(row['status'] == 'ERROR' for row in batch.results) else 'SUCCEEDED'
            )
            batch.save(update_fields=['state'])
    except AppError as error:
        batch.refresh_from_db()
        batch.state = 'FAILED'
        batch.results.append(
            {'status': 'ERROR', 'code': error.code, 'errors': {'detail': error.message}}
        )
        batch.save()
        raise
    record_audit(
        action='USER_BATCH_FINISHED',
        actor=batch.owner,
        target_type='account_batch',
        target_id=str(batch.pk),
        details={'kind': batch.kind, 'state': batch.state, 'count': len(batch.results)},
    )
    return {'batch_id': str(batch.pk), 'state': batch.state}


def process_rows(batch, completed, required):
    for row in batch.rows:
        if row['row'] in completed:
            continue
        try:
            with transaction.atomic():
                serialize_admin_changes()
                require_admin(get_user_model().objects.get(pk=batch.owner_id), required)
                if batch.expires_at <= timezone.now():
                    raise AppError(ErrorCode.CONFLICT, '批次执行已超时，请重新预校验')
                if row['status'] == 'ERROR':
                    result = {
                        key: value
                        for key, value in row.items()
                        if key not in {'input', 'snapshot', 'roles_snapshot'}
                    }
                else:
                    user, serializer = prepare(row, batch.mode)
                    if snapshot(user) != row.get('snapshot'):
                        raise AppError(ErrorCode.CONFLICT, '账号已被其他操作修改，请重新预校验')
                    if role_snapshot(serializer, user) != row.get('roles_snapshot'):
                        raise AppError(ErrorCode.CONFLICT, '角色已变更，请重新预校验')
                    saved = serializer.save()
                    record_audit(
                        action='USER_BATCH_ROW',
                        actor=batch.owner,
                        target_type='user',
                        target_id=str(saved.pk),
                        details={
                            'batch_id': str(batch.pk),
                            'row': row['row'],
                            'operation': row['status'],
                            'roles': sorted(saved.groups.values_list('name', flat=True)),
                            'node_scope': getattr(
                                getattr(saved, 'profile', None), 'node_scope', None
                            ),
                            'quota_overrides': getattr(
                                getattr(saved, 'profile', None), 'quota_overrides', {}
                            ),
                        },
                    )
                    result = {
                        'row': row['row'],
                        'username': saved.username,
                        'status': 'SUCCESS',
                        'operation': row['status'],
                    }
                batch.results.append(result)
                batch.save(update_fields=['results'])
        except (AppError, serializers.ValidationError, IntegrityError) as error:
            if batch.strict:
                raise AppError(ErrorCode.CONFLICT, '严格模式执行失败，整批已回滚') from error
            batch.refresh_from_db()
            result = (
                error_result(row, error)
                if not isinstance(error, IntegrityError)
                else {
                    'row': row['row'],
                    'username': row['input'].get('username', ''),
                    'status': 'ERROR',
                    'code': 'CONFLICT',
                    'errors': {'detail': '账号数据冲突，请重新预校验'},
                }
            )
            batch.results.append(result)
            batch.save(update_fields=['results'])


class AccountBatchView(APIView):
    permission_classes = [IsAuthenticated, RequiresPasswordChange]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    @transaction.atomic
    def post(self, request):
        kind = request.data.get('kind', 'import')
        require_admin(
            request.user, 'accounts.import_user' if kind == 'import' else 'accounts.export_user'
        )
        if kind not in ('import', 'export'):
            raise AppError(ErrorCode.INVALID_ARGUMENT, '批次类型无效')
        if (
            AccountBatch.objects.filter(
                owner=request.user,
                state__in=['DRAFT', 'PENDING', 'RUNNING'],
                expires_at__gt=timezone.now(),
            ).count()
            >= 5
        ):
            raise AppError(ErrorCode.CONFLICT, '同时最多保留 5 个未完成批次')
        # Lazily erase expired payloads; metadata and audit remain available.
        AccountBatch.objects.filter(expires_at__lte=timezone.now()).exclude(state='RUNNING').update(
            rows=[], results=[], file_data=None
        )
        if kind == 'import':
            mode = request.data.get('mode', 'create')
            if mode not in ('create', 'update', 'upsert'):
                raise AppError(ErrorCode.INVALID_ARGUMENT, '导入模式无效')
            upload = request.FILES.get('file')
            rows = parse_file(upload)
            batch = AccountBatch.objects.create(
                owner=request.user,
                kind=kind,
                mode=mode,
                format='xlsx' if upload.name.lower().endswith('.xlsx') else 'csv',
                strict=request.data.get('strict') is True or request.data.get('strict') == 'true',
                rows=preview(rows, mode),
                expires_at=timezone.now() + timedelta(minutes=30),
            )
        else:
            format = request.data.get('format', 'csv')
            fields = request.data.get(
                'fields', [field for field in EXPORT_FIELDS if field not in SENSITIVE]
            )
            if (
                format not in ('csv', 'xlsx')
                or not isinstance(fields, list)
                or not fields
                or len(fields) > len(EXPORT_FIELDS)
                or any(not isinstance(field, str) for field in fields)
                or len(set(fields)) != len(fields)
                or set(fields) - set(EXPORT_FIELDS)
            ):
                raise AppError(ErrorCode.INVALID_ARGUMENT, '导出格式或字段无效')
            if SENSITIVE & set(fields):
                require_admin(request.user, 'accounts.export_sensitive_user')
            query = get_user_model().objects.all()
            filters = request.data.get('filters', {})
            if not isinstance(filters, dict) or set(filters) - {'search', 'role', 'active'}:
                raise AppError(ErrorCode.INVALID_ARGUMENT, '导出筛选条件无效')
            search = str(filters.get('search', ''))[:150]
            if search:
                query = query.filter(
                    Q(username__icontains=search)
                    | Q(profile__display_name__icontains=search)
                    | Q(email__icontains=search)
                )
            if filters.get('role'):
                query = query.filter(groups__name=str(filters['role'])[:150])
            if filters.get('active') in ['true', 'false']:
                query = query.filter(is_active=filters['active'] == 'true')
            query = query.distinct()
            if request.data.get('ids') is not None:
                ids = request.data['ids']
                if (
                    not isinstance(ids, list)
                    or len(ids) > 10000
                    or any(type(value) is not int for value in ids)
                ):
                    raise AppError(ErrorCode.INVALID_ARGUMENT, '用户选择范围无效')
                query = query.filter(pk__in=ids)
                if query.count() != len(set(ids)):
                    raise AppError(ErrorCode.CONFLICT, '部分所选账号不存在或不再匹配筛选条件')
            if request.data.get('search'):
                query = query.filter(username__icontains=str(request.data['search'])[:150])
            if query.count() > 10000:
                raise AppError(
                    ErrorCode.INVALID_ARGUMENT, '每次最多导出 10000 个用户，请缩小筛选范围'
                )
            batch = AccountBatch.objects.create(
                owner=request.user,
                kind=kind,
                format=format,
                options={'fields': fields, 'ids': list(query.values_list('id', flat=True))},
                expires_at=timezone.now() + timedelta(minutes=30),
            )
            queue(batch)
        record_audit(
            action='USER_BATCH_CREATED',
            actor=request.user,
            target_type='account_batch',
            target_id=str(batch.pk),
            details={
                'kind': kind,
                'count': len(batch.rows) if kind == 'import' else len(batch.options['ids']),
            },
        )
        return Response(describe(batch), status=201 if kind == 'import' else 202)


def queue(batch):
    batch.task = enqueue(
        task_type='ACCOUNT_BATCH',
        resource_type='account_batch',
        resource_id=str(batch.pk),
        requested_by=batch.owner,
        payload={'batch_id': str(batch.pk)},
        idempotency_key=f'account-batch-{batch.pk}',
        max_attempts=1,
    )
    batch.state = 'PENDING'
    batch.save(update_fields=['task', 'state'])


class AccountBatchDetailView(APIView):
    permission_classes = [IsAuthenticated, RequiresPasswordChange]

    def owned(self, request, pk, lock=False):
        query = (
            AccountBatch.objects.select_for_update(of=('self',)) if lock else AccountBatch.objects
        )
        batch = get_object_or_404(query.select_related('task'), pk=pk, owner=request.user)
        require_admin(
            request.user,
            'accounts.import_user' if batch.kind == 'import' else 'accounts.export_user',
        )
        return batch

    def get(self, request, pk):
        batch = self.owned(request, pk)
        if batch.expires_at <= timezone.now():
            batch.rows = []
            batch.results = []
            batch.file_data = None
            batch.state = 'EXPIRED'
        response = Response(describe(batch))
        response['Cache-Control'] = 'no-store'
        return response

    @transaction.atomic
    def post(self, request, pk):
        batch = self.owned(request, pk, True)
        if batch.expires_at <= timezone.now():
            raise AppError(ErrorCode.CONFLICT, '预校验已过期')
        if batch.task_id:
            return Response(describe(batch), status=202)
        if batch.kind != 'import' or batch.state != 'DRAFT':
            raise AppError(ErrorCode.CONFLICT, '此批次无法提交')
        if batch.strict and any(row['status'] == 'ERROR' for row in batch.rows):
            raise AppError(ErrorCode.CONFLICT, '严格模式存在错误，请修正后重新上传')
        if not any(row['status'] != 'ERROR' for row in batch.rows):
            raise AppError(ErrorCode.CONFLICT, '没有可执行的有效记录')
        queue(batch)
        return Response(describe(batch), status=202)

    def delete(self, request, pk):
        batch = self.owned(request, pk)
        if batch.task_id and batch.task.status not in {'SUCCEEDED', 'FAILED', 'CANCELLED'}:
            raise AppError(ErrorCode.CONFLICT, '请等待后台任务结束')
        batch.rows = []
        batch.results = []
        batch.file_data = None
        batch.state = 'CANCELLED'
        batch.save()
        return Response(status=204)


class AccountBatchDownloadView(AccountBatchDetailView):
    def get(self, request, pk):
        batch = self.owned(request, pk)
        if batch.expires_at <= timezone.now():
            raise AppError(ErrorCode.CONFLICT, '下载已过期')
        if batch.kind == 'export':
            if SENSITIVE & set(batch.options['fields']):
                require_admin(request.user, 'accounts.export_sensitive_user')
            if not batch.file_data or batch.state != 'SUCCEEDED':
                raise AppError(ErrorCode.CONFLICT, '导出任务尚未成功完成')
            content = bytes(batch.file_data)
            format = batch.format
        else:
            rows = [
                {
                    'row': entry.get('row', ''),
                    'username': entry.get('username', ''),
                    'status': entry['status'],
                    'code': entry.get('code', ''),
                    'errors': json.dumps(entry.get('errors', {}), ensure_ascii=False),
                }
                for entry in (batch.results or batch.rows)
            ]
            content = file_bytes(['row', 'username', 'status', 'code', 'errors'], rows, 'csv')
            format = 'csv'
        record_audit(
            action='USER_BATCH_DOWNLOADED',
            actor=request.user,
            target_type='account_batch',
            target_id=str(pk),
            details={'kind': batch.kind, 'format': format},
        )
        return download(content, format, 'users-' + str(pk))


class AccountTemplateView(APIView):
    permission_classes = [IsAuthenticated, RequiresPasswordChange]

    def get(self, request):
        require_admin(request.user, 'accounts.import_user')
        format = request.query_params.get('file_format', 'csv')
        if format not in {'csv', 'xlsx'}:
            raise AppError(ErrorCode.INVALID_ARGUMENT, '格式仅支持 CSV 或 XLSX')
        return download(
            file_bytes(
                IMPORT_FIELDS,
                [
                    {
                        'username': 'user_example',
                        'display_name': '示例账号',
                        'role_codes': 'Viewer',
                        'status': 'pending',
                    }
                ],
                format,
            ),
            format,
            'users-template',
        )
